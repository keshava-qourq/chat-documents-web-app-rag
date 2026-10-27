"""Unit tests for `app/extraction.py` in isolation from the rest of the
ingestion pipeline (chunking, embeddings and persistence are covered in
`test_ingestion.py`, which also exercises `extract_fragments` for the four
supported extensions at a basic level -- one real PDF with no text layer,
one DOCX with and without headings, a plain TXT/MD file, an empty file, and
an unsupported extension). This file reaches the cases that coverage leaves
open: PDF page numbering across multiple pages (including a mix of blank and
non-blank pages, and `extract_text()` returning `None` rather than `""`),
DOCX paragraphs appearing before any heading, a heading whose own text is
blank, consecutive headings with no body between them, an empty DOCX
document, extension case-insensitivity, and plain-text content that is not
valid UTF-8.

No acceptance criteria are attached to this ticket ("Unit tests for
backend/app/extraction.py"); the cases below follow from the module's own
documented contract instead. The PDF tests fake `PdfReader` itself rather
than building a real text-layer PDF (which pypdf has no convenient way to
produce) -- what is under test here is `extraction._extract_pdf`'s own
page-numbering and blank-page-skipping logic, not pypdf's text extraction,
so a fake page object with a controllable `extract_text()` return value
exercises exactly that logic.
"""

import pytest
from docx import Document as DocxDocument

from app import extraction


class _FakePdfPage:
    def __init__(self, text: str | None) -> None:
        self._text = text

    def extract_text(self) -> str | None:
        return self._text


class _FakePdfReader:
    def __init__(self, pages_text: list[str | None]) -> None:
        self.pages = [_FakePdfPage(text) for text in pages_text]


def _patch_pdf_reader(monkeypatch, pages_text: list[str | None]) -> None:
    monkeypatch.setattr(extraction, "PdfReader", lambda stream: _FakePdfReader(pages_text))


# --- PDF ---------------------------------------------------------------


def test_extract_pdf_numbers_fragments_by_their_actual_one_based_page_number(monkeypatch) -> None:
    """A blank page in the middle must not shift the page number reported
    for the page after it -- page 3's fragment must still say "page 3", not
    "page 2"."""
    _patch_pdf_reader(monkeypatch, ["Page one text", "", "Page three text"])

    fragments = extraction.extract_fragments(b"irrelevant", ".pdf", "doc.pdf")

    assert fragments == [
        ("Page one text", "doc.pdf, page 1"),
        ("Page three text", "doc.pdf, page 3"),
    ]


def test_extract_pdf_extract_text_returning_none_is_treated_as_a_blank_page(monkeypatch) -> None:
    """pypdf's `extract_text()` can return `None` as well as `""`; both must
    be treated as "no text on this page", not raise on `.strip()`."""
    _patch_pdf_reader(monkeypatch, [None, "Real content"])

    fragments = extraction.extract_fragments(b"irrelevant", ".pdf", "doc.pdf")

    assert fragments == [("Real content", "doc.pdf, page 2")]


def test_extract_pdf_whitespace_only_page_text_yields_no_fragment(monkeypatch) -> None:
    _patch_pdf_reader(monkeypatch, ["   \n\t  "])

    fragments = extraction.extract_fragments(b"irrelevant", ".pdf", "doc.pdf")

    assert fragments == []


def test_extract_pdf_strips_surrounding_whitespace_from_page_text(monkeypatch) -> None:
    _patch_pdf_reader(monkeypatch, ["  padded text  \n"])

    fragments = extraction.extract_fragments(b"irrelevant", ".pdf", "doc.pdf")

    assert fragments == [("padded text", "doc.pdf, page 1")]


# --- DOCX ---------------------------------------------------------------


def _docx_bytes(document: DocxDocument) -> bytes:
    import io

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_docx_text_before_the_first_heading_is_located_at_the_file_name() -> None:
    """Paragraphs appearing before any heading have no heading to be grouped
    under, so their fragment's location must fall back to the file name --
    and a heading appearing afterwards must switch later paragraphs to being
    located under it."""
    document = DocxDocument()
    document.add_paragraph("Preamble text with no heading above it.")
    document.add_heading("Later Section", level=1)
    document.add_paragraph("Body under the later section.")

    fragments = extraction.extract_fragments(_docx_bytes(document), ".docx", "report.docx")

    assert fragments == [
        ("Preamble text with no heading above it.", "report.docx"),
        ("Body under the later section.", "report.docx, Later Section"),
    ]


def test_extract_docx_consecutive_headings_with_no_body_produce_no_fragment_for_the_empty_one() -> (
    None
):
    document = DocxDocument()
    document.add_heading("Empty Section", level=1)
    document.add_heading("Section With Content", level=1)
    document.add_paragraph("Only this section has a body.")

    fragments = extraction.extract_fragments(_docx_bytes(document), ".docx", "doc.docx")

    assert fragments == [("Only this section has a body.", "doc.docx, Section With Content")]


def test_extract_docx_heading_with_blank_text_keeps_the_previous_heading_as_location() -> None:
    """An empty/whitespace-only heading paragraph must not blank out the
    location for the body that follows it -- the code falls back to keeping
    the previous heading rather than switching to `None`."""
    document = DocxDocument()
    document.add_heading("Real Heading", level=1)
    document.add_heading("   ", level=2)
    document.add_paragraph("Body after the blank heading.")

    fragments = extraction.extract_fragments(_docx_bytes(document), ".docx", "doc.docx")

    assert fragments == [("Body after the blank heading.", "doc.docx, Real Heading")]


def test_extract_docx_empty_document_returns_no_fragments() -> None:
    document = DocxDocument()

    fragments = extraction.extract_fragments(_docx_bytes(document), ".docx", "empty.docx")

    assert fragments == []


def test_extract_docx_multiple_paragraphs_under_one_heading_are_joined_into_one_fragment() -> None:
    document = DocxDocument()
    document.add_heading("Section", level=1)
    document.add_paragraph("First line.")
    document.add_paragraph("Second line.")

    fragments = extraction.extract_fragments(_docx_bytes(document), ".docx", "doc.docx")

    assert fragments == [("First line.\nSecond line.", "doc.docx, Section")]


# --- dispatch / plain text ------------------------------------------------


def test_extract_fragments_extension_matching_is_case_insensitive() -> None:
    fragments = extraction.extract_fragments(b"hello", ".TXT", "readme.txt")
    assert fragments == [("hello", "readme.txt")]


def test_extract_plain_invalid_utf8_bytes_are_replaced_rather_than_raising() -> None:
    """Non-UTF-8 bytes must not raise `UnicodeDecodeError`: the module
    docstring commits to extraction always producing a terminal document
    status (ready or failed with a clear reason), never an unhandled
    exception from a byte sequence a user's editor happened to save."""
    content = b"before \xff\xfe after"

    fragments = extraction.extract_fragments(content, ".txt", "garbled.txt")

    assert len(fragments) == 1
    text, location = fragments[0]
    assert location == "garbled.txt"
    assert "before" in text
    assert "after" in text


def test_extract_fragments_unsupported_extension_message_names_the_extension() -> None:
    with pytest.raises(ValueError, match=r"\.xlsx"):
        extraction.extract_fragments(b"data", ".xlsx", "sheet.xlsx")
